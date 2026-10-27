# -*- coding: utf-8 -*-
import openpyxl, sys
from openpyxl.styles import Font, PatternFill

sys.stdout.reconfigure(encoding='utf-8')

excel_paths = [
    "Testing/TestCases_KLCN_CoTuong (1).xlsx",
    "Testing/04_TaiLieu_Va_BaoCao_KiemThu/BANG_KE_HOACH_VA_KICH_BAN_KIEM_THU_137_TC.xlsx"
]

FILL_FAIL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
FONT_FAIL = Font(name="Segoe UI", size=10, bold=True, color="990000")

# 7 Real Defects detected through multi-tool probing
real_bugs = [
    {
        "stt": 1,
        "bug_id": "BUG_001",
        "title": "Cho phép đăng ký mật khẩu yếu dưới 6 ký tự qua API",
        "module": "Xác thực & Tài khoản (04_TC_Auth_Profile)",
        "severity": "P2 - Major",
        "priority": "High",
        "steps": "1. Gửi POST /api/auth/register\n2. Body chứa password = '123'\n3. Kiểm tra mã phản hồi",
        "expected": "HTTP 400 Bad Request, báo lỗi mật khẩu phải >= 6 ký tự",
        "actual": "HTTP 200 OK, hệ thống tạo tài khoản thành công với mật khẩu 3 ký tự",
        "screenshot": "E2E_SCREENSHOTS/TC_AUTH_004.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "04_TC_Auth_Profile",
        "tc_code": "TC_AUTH_004"
    },
    {
        "stt": 2,
        "bug_id": "BUG_002",
        "title": "Thiếu Regex kiểm tra định dạng Email khi đăng ký tài khoản",
        "module": "Xác thực & Tài khoản (04_TC_Auth_Profile)",
        "severity": "P2 - Major",
        "priority": "Medium",
        "steps": "1. Gửi POST /api/auth/register\n2. Body email = 'not_an_email_at_all' (không có @ và domain)\n3. Bấm đăng ký",
        "expected": "HTTP 400 Bad Request, từ chối định dạng email không hợp lệ",
        "actual": "HTTP 200 OK, tài khoản vẫn được tạo với email rác trong CSDL",
        "screenshot": "E2E_SCREENSHOTS/TC_AUTH_001.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "04_TC_Auth_Profile",
        "tc_code": "TC_AUTH_001"
    },
    {
        "stt": 3,
        "bug_id": "BUG_003",
        "title": "Truy cập /Admin khi chưa phân quyền trả về 302 Redirect thay vì 403 Forbidden",
        "module": "Phân hệ Quản trị viên (10_TC_Admin_System)",
        "severity": "P2 - Major",
        "priority": "High",
        "steps": "1. Người dùng chưa đăng nhập hoặc role User truy cập URL /Admin\n2. Kiểm tra HTTP Status Code trả về",
        "expected": "HTTP 403 Forbidden (hoặc trang Access Denied 403)",
        "actual": "HTTP 302 Found, chuyển hướng (Redirect) về /Auth/Login?ReturnUrl=%2FAdmin",
        "screenshot": "E2E_SCREENSHOTS/TC_ADM_002.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "10_TC_Admin_System",
        "tc_code": "TC_ADM_002"
    },
    {
        "stt": 4,
        "bug_id": "BUG_004",
        "title": "Vẫn cho phép đi quân cờ sau khi người chơi đã bấm Xin hàng (Surrender)",
        "module": "Đấu trực tuyến & Vận hành ván cờ (07_TC_Online_SignalR)",
        "severity": "P1 - Blocker",
        "priority": "High",
        "steps": "1. Tạo ván cờ mới\n2. Bên Đỏ gọi POST /api/games/{id}/surrender\n3. Bên Đỏ tiếp tục gửi nước đi POST /api/games/{id}/moves",
        "expected": "HTTP 400 Bad Request, ván cờ đã kết thúc không thể đi thêm nước cờ",
        "actual": "HTTP 200 OK, server tiếp tục nhận nước cờ và đổi trạng thái về InProgress",
        "screenshot": "E2E_SCREENSHOTS/TC_ONL_015.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "07_TC_Online_SignalR",
        "tc_code": "TC_ONL_015"
    },
    {
        "stt": 5,
        "bug_id": "BUG_005",
        "title": "Cấp Chứng chỉ tốt nghiệp cho người dùng chưa hoàn thành khóa học (0% tiến độ)",
        "module": "Học viện & Kỳ phổ (09_TC_Course_Master)",
        "severity": "P2 - Major",
        "priority": "High",
        "steps": "1. Người dùng mới tạo chưa học bài nào\n2. Truy cập trực tiếp /Course/Certificate?courseId=1\n3. Kiểm tra trang hiển thị",
        "expected": "Từ chối cấp chứng chỉ, thông báo học viên cần hoàn thành 100% khóa học",
        "actual": "HTTP 200 OK, hệ thống vẫn kết xuất chứng chỉ tốt nghiệp hoàn chỉnh",
        "screenshot": "E2E_SCREENSHOTS/TC_COU_009.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "09_TC_Course_Master",
        "tc_code": "TC_COU_009"
    },
    {
        "stt": 6,
        "bug_id": "BUG_006",
        "title": "API tiến độ khóa học không chặn ngưỡng điểm Quiz (Chấp nhận điểm 999999%)",
        "module": "Học viện & Kỳ phổ (09_TC_Course_Master)",
        "severity": "P3 - Minor",
        "priority": "Medium",
        "steps": "1. Gửi POST /api/courses/progress với quizScore = 999999\n2. Kiểm tra phản hồi máy chủ",
        "expected": "HTTP 400 Bad Request, điểm số Quiz phải nằm trong khoảng 0 đến 100%",
        "actual": "HTTP 200 OK, máy chủ lưu nguyên vẹn bestQuizScore = 999999",
        "screenshot": "E2E_SCREENSHOTS/TC_COU_008.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "09_TC_Course_Master",
        "tc_code": "TC_COU_008"
    },
    {
        "stt": 7,
        "bug_id": "BUG_007",
        "title": "Chức năng Xóa bài cờ thế trong Admin trả về HTTP 200 khi ID không tồn tại",
        "module": "Phân hệ Quản trị viên (10_TC_Admin_System)",
        "severity": "P3 - Minor",
        "priority": "Low",
        "steps": "1. Gửi POST /Admin/DeletePuzzle với id = 999999 không có trong CSDL\n2. Kiểm tra mã phản hồi",
        "expected": "HTTP 404 Not Found hoặc báo lỗi không tìm thấy bài cờ cần xóa",
        "actual": "HTTP 200 OK (hoặc redirect thành công) gây sai lệch trạng thái CMS",
        "screenshot": "E2E_SCREENSHOTS/TC_ADM_009.png",
        "tester": "Tester độc lập",
        "status": "Open",
        "sheet": "10_TC_Admin_System",
        "tc_code": "TC_ADM_009"
    }
]

for path in excel_paths:
    print(f"Updating bugs into: {path}")
    wb = openpyxl.load_workbook(path)
    
    # 1. Update Sheet 12: Defect Log
    ws12 = wb["12_Defect_Log"]
    for b in real_bugs:
        r = 4 + b["stt"]
        ws12.cell(row=r, column=1, value=b["stt"])
        ws12.cell(row=r, column=2, value=b["bug_id"])
        ws12.cell(row=r, column=3, value=b["title"])
        ws12.cell(row=r, column=4, value=b["module"])
        ws12.cell(row=r, column=5, value=b["severity"])
        ws12.cell(row=r, column=6, value=b["priority"])
        ws12.cell(row=r, column=7, value=b["steps"])
        ws12.cell(row=r, column=8, value=b["expected"])
        ws12.cell(row=r, column=9, value=b["actual"])
        ws12.cell(row=r, column=10, value=b["screenshot"])
        ws12.cell(row=r, column=11, value=b["tester"])
        c_st = ws12.cell(row=r, column=12, value=b["status"])
        c_st.font = FONT_FAIL
        c_st.fill = FILL_FAIL
    
    # 2. Update status of corresponding test cases to 'Fail'
    for b in real_bugs:
        sheet_name = b["sheet"]
        tc_code = b["tc_code"]
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in range(5, ws.max_row + 1):
                if str(ws.cell(row=row, column=2).value or "").strip() == tc_code:
                    ws.cell(row=row, column=9, value=b["actual"])
                    c_status = ws.cell(row=row, column=10, value="Fail")
                    c_status.font = FONT_FAIL
                    c_status.fill = FILL_FAIL
                    ws.cell(row=row, column=12, value=b["bug_id"])
                    break
    
    wb.save(path)
    print(f"Saved: {path}")

print("All 7 bugs recorded successfully into both Excel files!")
